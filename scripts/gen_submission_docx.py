"""Generate SauceDemo-AI-QA-Submission.docx with embedded evidence."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "evidence"
OUT = ROOT / "SauceDemo-AI-QA-Submission.docx"

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)


def font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        font(r, size={1: 16, 2: 13, 3: 12}.get(level, 12), bold=True)
    return p


def p(text, bold=False, italic=False, size=11, after=6):
    para = doc.add_paragraph()
    run = para.add_run(text)
    font(run, size=size, bold=bold)
    run.italic = italic
    para.paragraph_format.space_after = Pt(after)
    return para


def bullet(text):
    para = doc.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    font(run, size=10)
    para.paragraph_format.space_after = Pt(3)


def shade_header(row):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
                run.font.size = Pt(9)
                run.font.name = "Calibri"


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
        for para in t.rows[0].cells[i].paragraphs:
            for run in para.runs:
                font(run, size=9, bold=True)
    shade_header(t.rows[0])
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            t.rows[r_i + 1].cells[c_i].text = val
            for para in t.rows[r_i + 1].cells[c_i].paragraphs:
                for run in para.runs:
                    font(run, size=9)
    doc.add_paragraph()


def img(path: Path, width=5.8, caption=None):
    if not path.exists():
        p(f"[Missing evidence: {path.name}]", italic=True, size=9)
        return
    doc.add_picture(str(path), width=Inches(width))
    if caption:
        cap = doc.add_paragraph()
        run = cap.add_run(caption)
        font(run, size=9, bold=False)
        run.italic = True
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)


# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SauceDemo — AI-Assisted QA Submission")
font(r, size=20, bold=True, color=(31, 78, 121))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Tasks A–D · Defect Report · Automation · AI QA Strategy")
font(r, size=12, bold=True)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run(
    "App: https://www.saucedemo.com/  |  Account: standard_user / secret_sauce\n"
    "Date: 6 October 2026  |  Automation: Playwright (TypeScript)"
)
font(r, size=10)
meta.paragraph_format.space_after = Pt(12)

h("1. Executive summary")
p(
    "AI accelerated risk discovery and automation drafting for SauceDemo. Ideas were challenged "
    "against the live app on standard_user. Five high-impact defects were confirmed; unsupported "
    "AI suggestions (real payments, stock, i18n, etc.) were discarded. One maintainable Playwright "
    "E2E covers the critical purchase path."
)

# ---------------------------------------------------------------------------
# Task A — full exploratory testing section (before defects)
# ---------------------------------------------------------------------------
h("2. Task A — AI-Assisted Exploratory Testing")
p(
    "Goal: use AI to accelerate exploration of SauceDemo — not to dump a traditional test-case list. "
    "AI proposed risks and edge cases; each idea was converted into a live check, then kept, changed, "
    "or discarded based on what the app actually does for standard_user."
)

h("2.1 Approach", 2)
bullet("Explore live app at https://www.saucedemo.com/ with standard_user / secret_sauce.")
bullet("Ask AI for high-risk journeys, functional risks, edge cases, UX/a11y, performance, and follow-up questions.")
bullet("Execute useful suggestions as real tests in the browser.")
bullet("Challenge or discard ideas that are irrelevant, wrong, redundant, or unsupported by the app.")

h("2.2 Highest-risk user journeys identified", 2)
table(
    ["Journey", "Why it mattered", "Session outcome"],
    [
        [
            "Purchase / checkout integrity",
            "Money, order completion, and validation sit here",
            "Highest risk — multiple critical defects (see Task B)",
        ],
        [
            "Cart state (add/remove/empty/badge)",
            "Empty-cart checkout and no qty control change purchase semantics",
            "High risk — empty checkout confirmed",
        ],
        [
            "Dynamic Catalog (Lazy / Spinner / Slider)",
            "Newer surface with loading, duplicates, a11y, browse-only UX",
            "High risk — duplicate SKUs confirmed",
        ],
        [
            "Auth + session guards",
            "Protects all post-login pages",
            "Medium — mostly passed (guards + locked_out_user)",
        ],
        [
            "Catalog browse + sort",
            "Core discovery path",
            "Lower risk — sort/tax healthy for standard_user",
        ],
    ],
)

h("2.3 What was actually tested (live)", 2)
bullet("Login: empty submit, username-only, password-only, locked_out_user.")
bullet("Inventory: price sort low→high and high→low; add/remove; cart badge.")
bullet("Cart: quantity editability, localStorage cart model, empty-cart checkout.")
bullet("Checkout: whitespace-only names/zip, invalid zip (!!!), tax math, Finish, Back then re-Finish.")
bullet("Order complete: thank-you page; Generate PDF order button state.")
bullet("Auth guards: logout then direct URL to /inventory.html and /checkout-step-one.html.")
bullet("Dynamic Catalog: Lazy Load scroll/duplicates, Spinner page, Slider dots/a11y.")

h("2.4 Exploration outcomes (summary)", 2)
p("Confirmed bugs (detail + evidence in Task B):", bold=True, after=2)
bullet("Empty cart can complete a $0 order.")
bullet("Whitespace-only First/Last/Zip accepted.")
bullet("Postal code accepts “!!!” with no format check.")
bullet("Browser Back after Thank You still allows Finish.")
bullet("Lazy Load catalog duplicates Test.allTheThings() T-Shirt (Red) (L).")
p("Confirmed passes:", bold=True, after=2)
bullet("Tax math at 8% correct ($29.99 → tax $2.40 → total $32.39; multi-item $79.98 → $6.40 → $86.38).")
bullet("Price sort lohi/hilo correct; cart badge increments on Add.")
bullet("Route guards after logout; locked_out_user blocked.")
p("Limitations (not filed as defects):", bold=True, after=2)
bullet("Dynamic Catalog pages are browse-only (no Add to cart).")
bullet("Cart quantity cannot exceed 1 (localStorage stores product ID array only).")

h("2.5 Representative AI prompts (Task A deliverable)", 2)
p(
    "Five prompts used to accelerate exploration. For each: purpose, useful AI output, and what was "
    "tested / changed / rejected / validated against the live app."
)

prompts_a = [
    (
        "Prompt 1 — Risk prioritization",
        "“For SauceDemo (saucedemo.com), identify the highest-risk user journeys for standard_user. "
        "Prioritize by business/impact risk, not by page count. Challenge assumptions that only apply "
        "to real e-commerce backends.”",
        "Focus the session on risk instead of generating a generic test-case list.",
        "Ranked checkout/cart as top risk; auth/session next; warned payment may be mocked.",
        "Validated checkout + cart as highest risk in the live app. Kept auth as medium after "
        "logout/direct-URL checks. Rejected deep payment/fraud testing after confirming SauceCard is fake.",
    ),
    (
        "Prompt 2 — Checkout edge cases",
        "“Generate functional risks and edge cases for SauceDemo checkout (info form → overview → finish). "
        "Include empty cart, whitespace, postal formats, back-button after complete, and tax math. "
        "Mark which need live proof.”",
        "Convert AI edge-case ideation into concrete checkout experiments.",
        "Suggested empty-cart checkout, whitespace names, invalid zip, tax rounding, re-submit via Back.",
        "Validated as bugs: empty cart $0 order, whitespace accepted, zip “!!!” accepted, Back+Finish "
        "replay. Validated as pass: 8% tax. Tax idea confirmed correct — no change needed.",
    ),
    (
        "Prompt 3 — Dynamic Catalog scout",
        "“After login, the menu has Dynamic Catalog → Lazy Load / Spinner / Slider. What should I explore "
        "for data integrity, loading UX, a11y, and purchase flow gaps? Discard ideas if those pages are "
        "browse-only.”",
        "Investigate a less obvious, higher-novelty surface with AI as a scout.",
        "Flagged infinite-scroll duplicates, spinner leftovers, slider keyboard/a11y, missing Add to cart.",
        "Validated duplicate Red T-Shirt (L); all slider dots a11y-current; browse-only. Treated missing "
        "Add to cart as limitation, not defect. Keyboard ArrowLeft did not move slider — kept as open follow-up.",
    ),
    (
        "Prompt 4 — Login / session risks (scoped)",
        "“Propose login/UX and session risks for SauceDemo: empty fields, partial credentials, "
        "locked_out_user, and direct URL access after logout. Ignore SQL injection unless the UI gives "
        "a real reason to prioritize it.”",
        "Cover auth without wasting time on generic security folklore.",
        "Required-field messaging, locked-out persona, route guards, misleading error chrome.",
        "Validated required-field errors, locked_out_user blocked, route guards. Confirmed UX issue: both "
        "fields show red X when only one is wrong. Rejected SQL-injection deep dive as out of scope.",
    ),
    (
        "Prompt 5 — Explicit discard pass",
        "“Which of these AI ideas should I discard for SauceDemo and why: multi-currency, stock limits, "
        "email confirmation, multi-user cart sync, real card processing? Base the answer on what the app "
        "actually exposes.”",
        "Force explicit rejection of AI noise before writing findings.",
        "Discard list with rationale tied to missing UI/capabilities.",
        "All five discarded after live observation (no locale/currency, no stock, no email at checkout, "
        "cart is localStorage ID array, payment mocked). Prevented false findings in the report.",
    ),
]
for title_t, prompt, purpose, useful, action in prompts_a:
    h(title_t, 3)
    p("Prompt:", bold=True, after=2)
    p(prompt, italic=True)
    p("Purpose: " + purpose)
    p("Useful AI output: " + useful)
    p("What I tested / changed / rejected / validated: " + action)

h("2.6 AI suggestions discarded", 2)
table(
    ["Suggestion", "Why discarded"],
    [
        [
            "Deep payment / fraud / card-processor tests",
            "Payment is mocked (“SauceCard #31337”); no real processor or card form",
        ],
        [
            "SQL injection / auth bypass fuzzing as primary focus",
            "Demo credential list + cookie session-username; not worth session budget",
        ],
        ["i18n / multi-currency checkout", "No locale or currency switcher found"],
        [
            "Inventory stock / oversell limits",
            "No stock counters or purchase limits in UI or cart model",
        ],
        [
            "Multi-user cart sync / concurrency",
            "Cart is localStorage product-ID list on the client only",
        ],
        [
            "Email / order confirmation delivery",
            "Checkout collects name + zip only; no email field",
        ],
    ],
)

h("3. Task B — Defect report (top 5)")

# DEF-01
h("DEF-01 — Empty cart can complete checkout ($0 order)", 2)
table(
    ["Field", "Detail"],
    [
        ["Severity / Priority", "Critical / P0"],
        [
            "Steps to reproduce",
            "1) Login as standard_user. 2) Open empty cart. 3) Click Checkout. "
            "4) Enter any First/Last/Zip. 5) Continue → Finish.",
        ],
        ["Expected", "Checkout blocked when cart is empty; order cannot complete."],
        [
            "Actual",
            "Checkout proceeds. Overview: Item total $0 / Tax $0.00 / Total $0.00. "
            "Finish shows “Thank you for your order!”",
        ],
        [
            "Business / user impact",
            "Meaningless successful orders; broken purchase invariant; pollutes metrics.",
        ],
        [
            "Classification reasoning",
            "Reproduced on standard_user. Not a persona quirk. Empty-cart Checkout is a rule failure.",
        ],
    ],
)
img(EVIDENCE / "DEF01-empty-cart-checkout-enabled.png", caption="Fig. DEF-01a — Empty cart with Checkout enabled")
img(EVIDENCE / "DEF01-empty-order-overview-zero-total.png", caption="Fig. DEF-01b — Overview totals $0.00")
img(EVIDENCE / "DEF01-empty-order-thank-you.png", caption="Fig. DEF-01c — Thank-you after empty order")

# DEF-02
h("DEF-02 — Whitespace-only checkout identity accepted", 2)
table(
    ["Field", "Detail"],
    [
        ["Severity / Priority", "Critical / P0"],
        [
            "Steps to reproduce",
            "1) Login, add item, Checkout. 2) Enter spaces only in First/Last/Zip. 3) Continue.",
        ],
        ["Expected", "Validation error after trim; spaces-only rejected."],
        ["Actual", "Advances to Overview with no error."],
        [
            "Business / user impact",
            "Orders with unusable customer identity; fulfillment would fail in a real system.",
        ],
        [
            "Classification reasoning",
            "Truly empty fields are validated (“is required”), so accepting spaces is incomplete validation — not “no validation by design.”",
        ],
    ],
)
img(EVIDENCE / "DEF02-whitespace-only-checkout-fields.png", caption="Fig. DEF-02a — Fields contain only spaces")
img(EVIDENCE / "DEF02-whitespace-accepted-overview.png", caption="Fig. DEF-02b — Overview reached after whitespace submit")

# DEF-03
h("DEF-03 — Postal code accepts invalid format (!!!)", 2)
table(
    ["Field", "Detail"],
    [
        ["Severity / Priority", "High / P1"],
        [
            "Steps to reproduce",
            "1) Login, add item, Checkout. 2) Valid names + Zip “!!!”. 3) Continue.",
        ],
        ["Expected", "Format validation rejects non-postal values."],
        ["Actual", "Accepted; Overview reached."],
        ["Business / user impact", "Invalid shipping addresses; support/return cost."],
        [
            "Classification reasoning",
            "Field is labeled Zip/Postal Code and other fields are validated — format check is reasonable.",
        ],
    ],
)
img(EVIDENCE / "DEF03-invalid-postal-code.png", caption="Fig. DEF-03 — Zip value “!!!” accepted on the form")

# DEF-04
h("DEF-04 — Browser Back after completion allows re-Finish", 2)
table(
    ["Field", "Detail"],
    [
        ["Severity / Priority", "High / P1"],
        [
            "Steps to reproduce",
            "1) Complete checkout to Thank You. 2) Browser Back. 3) Finish still available; click again.",
        ],
        ["Expected", "Overview invalidated; Finish cannot re-submit."],
        ["Actual", "Overview + Finish remain usable (duplicate completion path)."],
        ["Business / user impact", "Duplicate submissions; confused order state."],
        [
            "Classification reasoning",
            "Classic history/state bug on standard_user — not an exotic persona issue.",
        ],
    ],
)
img(EVIDENCE / "DEF04-back-allows-refinish.png", caption="Fig. DEF-04 — Back returns to Overview with Finish active")

# DEF-05
h("DEF-05 — Lazy Load catalog duplicate SKUs", 2)
table(
    ["Field", "Detail"],
    [
        ["Severity / Priority", "High / P1"],
        [
            "Steps to reproduce",
            "1) Login. 2) Menu → Dynamic Catalog → Lazy Load. 3) Scroll to load items. "
            "4) Observe Test.allTheThings() T-Shirt (Red) (L).",
        ],
        ["Expected", "Each size/SKU appears once."],
        ["Actual", "Duplicate cards for the same name/size (worsens with further lazy loads)."],
        [
            "Business / user impact",
            "Catalog data integrity; user confusion; analytics/automation uniqueness breaks.",
        ],
        [
            "Classification reasoning",
            "Main inventory has unique items; bug is specific to Dynamic Catalog Lazy Load. "
            "Browse-only (no Add to cart) is a limitation, not this defect.",
        ],
    ],
)
img(EVIDENCE / "DEF05-lazy-load-duplicate-sku.png", caption="Fig. DEF-05 — Duplicate Red T-Shirt (L) highlighted")

h("Not reported as defects (reasoning)", 2)
bullet("Dynamic Catalog has no Add to cart → feature limitation / separate surface.")
bullet("Cart qty cannot exceed 1 → demo localStorage ID-array model.")
bullet("SauceCard #31337 mock payment → intentional demo behavior.")
bullet("problem_user / visual_user quirks → persona-specific, not primary-account defects.")

h("4. Task C — Automated test")
p(
    "Framework: Playwright (TypeScript). Flow: Login → add Sauce Labs Backpack → cart → "
    "checkout info → overview (assert totals) → finish → assert thank-you + cleared cart badge."
)
h("Setup & run", 3)
p("Repository folder: saucedemo-ai-qa (public GitHub — see README after publish).", italic=True)
bullet("npm install")
bullet("npx playwright install chromium")
bullet("npm test   # or: npm run test:headed")
p("Credentials: standard_user / secret_sauce (public demo credentials).")

h("AI assistance vs manual changes", 3)
table(
    ["Topic", "Detail"],
    [
        [
            "AI assisted",
            "Playwright+TS choice; happy-path structure; data-test/role locators; price assertions; storage isolation idea.",
        ],
        [
            "AI incorrect / brittle",
            "Assumed data-testid (SauceDemo uses data-test) → 60s timeout; early draft used hard waits; nth button locators.",
        ],
        [
            "Manual fixes",
            "testIdAttribute: 'data-test'; removed hard waits; product-specific data-test IDs; exact tax/total asserts; clear cookies+storage in beforeEach.",
        ],
        [
            "Limitations / risks",
            "Depends on live SauceDemo; single persona; does not encode known demo defects as expected failures; network flakiness possible.",
        ],
    ],
)
p("Test file: tests/checkout-happy-path.spec.ts — verified locally: 1 passed (~3s).", bold=True)

h("5. Task D — AI-driven QA approach (1 page)")
p(
    "When developers use AI to generate code, tests, and docs, QA shifts from writing every case "
    "by hand to risk ownership, oracle definition, and gatekeeping."
)
p("Where AI adds value:", bold=True, after=2)
bullet("Planning: risk heatmaps, journey candidates, ambiguity questions.")
bullet("Design: boundaries/negatives, state-transition gaps.")
bullet("Authoring: draft automation, locator suggestions, fixtures.")
bullet("Execution: flake clustering, failure summarization, log/trace triage.")
bullet("Reporting: defect drafts, release risk summaries.")

p("Where human review is mandatory:", bold=True, after=2)
bullet("Acceptance criteria / product oracles (“what is correct?”).")
bullet("Severity & release-risk decisions.")
bullet("Security, privacy, payments, compliance-sensitive flows.")
bullet("Any AI claim not reproduced on a real build.")
bullet("Final sign-off that a finding is a real defect (not persona/env quirk).")

p("Reviewing AI-generated tests before regression:", bold=True, after=2)
bullet("Map each test to a risk/requirement; reject orphans.")
bullet("Assert outcomes (URL, totals, persistence, errors), not only clicks.")
bullet("Prefer data-test / roles; ban brittle nth/XPath unless justified.")
bullet("Prove isolation; ban unnecessary hard waits.")
bullet("Run locally + CI; human marks “promotion approved.”")

p("Defect analysis, coverage, regression optimization:", bold=True, after=2)
bullet("AI clusters failures; humans confirm root cause.")
bullet("AI maps changed files → impacted journeys for regression subsets.")
bullet("AI finds coverage gaps vs critical journeys; humans prioritize.")
bullet("Optimize: quarantine flakes; keep critical path green.")

p("CI/CD quality gates:", bold=True, after=2)
bullet("PR: lint + unit + smoke E2E (purchase-class).")
bullet("Nightly: broader regression; AI summarizes failures with owners.")
bullet("Release: critical-path pass required; AI drafts risk note from open defects.")
bullet("AI never auto-closes a failing gate without human override policy.")

p("Key risks & mitigations:", bold=True, after=2)
table(
    ["Risk", "Mitigation"],
    [
        ["Wrong oracle in AI tests", "Human AC review; tie tests to requirements"],
        ["Brittle selectors / env", "data-test contract; target-env CI runs"],
        ["Hallucinated defects", "Require live repro + evidence"],
        ["Security blind spots", "Mandatory human/security review on auth/PII"],
        ["Automation theater", "Track escaped defects; keep exploratory charters"],
        ["Secret leakage in prompts", "Redact secrets; use public demo accounts in examples"],
    ],
)
p(
    "Bottom line: AI accelerates ideation and drafting; humans own correctness, risk, and "
    "promotion into the regression suite.",
    bold=True,
)

h("6. Tools used")
bullet("Cursor Agent / Composer — ideation, coding, packaging")
bullet("Cursor IDE Browser — live exploratory validation")
bullet("Playwright + Chromium — E2E automation")
bullet("Node.js / npm")
bullet("python-docx — Word submission export")
bullet("SauceDemo public web application")

h("7. Evidence & repo index")
table(
    ["Artifact", "Location"],
    [
        ["Automated test", "tests/checkout-happy-path.spec.ts"],
        ["Run instructions", "README.md"],
        ["AI usage detail", "docs/AI-USAGE.md"],
        ["Markdown narrative", "docs/SUBMISSION.md"],
        ["Screenshots", "evidence/DEF01–DEF05*.png"],
        ["This Word document", "SauceDemo-AI-QA-Submission.docx"],
    ],
)

doc.save(OUT)
print(OUT)
